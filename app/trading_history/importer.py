"""Bounded CSV, XLSX and screenshot preview parsing for trade imports."""

from __future__ import annotations

import csv
from datetime import datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
import io
import re
from statistics import median
from typing import Any
from zoneinfo import ZoneInfo

from .contracts import TradeImportPreview, TradePreviewRow


MAX_TABLE_BYTES = 10 * 1024 * 1024
MAX_IMAGE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20000
MAX_IMAGES = 10
OCR_CONFIDENCE_THRESHOLD = Decimal("0.85")
SHANGHAI = ZoneInfo("Asia/Shanghai")


class ImportParseError(ValueError):
    pass


class ImportLimitError(ImportParseError):
    pass


FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "traded_at": ("成交时间", "交易时间", "成交日期", "日期", "委托时间", "tradedat", "datetime", "date"),
    "security_code": ("证券代码", "股票代码", "基金代码", "代码", "symbol", "code"),
    "security_name": ("证券名称", "股票名称", "基金名称", "名称", "name"),
    "side": ("买卖方向", "操作", "业务名称", "方向", "side"),
    "quantity": ("成交数量", "数量", "发生数量", "成交股数", "quantity", "qty"),
    "price_cny": ("成交价格", "成交均价", "价格", "price"),
    "gross_amount_cny": ("成交金额", "发生金额", "金额", "amount"),
    "fee_cny": ("手续费", "佣金", "费用", "fee"),
    "asset_type": ("证券类别", "资产类型", "品种", "assettype"),
    "account_alias": ("账户", "资金账号", "账户名称", "account"),
    "broker_trade_id": ("成交编号", "合同编号", "流水号", "tradeid", "orderid"),
    "account_value_cny": ("总资产", "账户总资产", "净资产", "accountvalue"),
    "currency": ("币种", "currency"),
}


def _key(value: object) -> str:
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]", "", str(value or "").strip().casefold())


def _mapping(columns: list[str]) -> dict[str, str]:
    result: dict[str, str] = {}
    normalized = {_key(column): column for column in columns}
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            candidate = normalized.get(_key(alias))
            if candidate is not None:
                result[field] = candidate
                break
    return result


def _decimal(value: object) -> Decimal | None:
    text = re.sub(r"[^0-9.\-]", "", str(value or "").replace(",", ""))
    if not text or text in {"-", ".", "-."}:
        return None
    try:
        parsed = Decimal(text)
    except InvalidOperation:
        return None
    return parsed if parsed.is_finite() else None


def _datetime(value: object) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=value.tzinfo or SHANGHAI)
    text = str(value or "").strip()
    if not text:
        return None
    normalized = text.replace("年", "-").replace("月", "-").replace("日", " ").replace("/", "-")
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y%m%d %H:%M:%S", "%Y%m%d"):
        try:
            return datetime.strptime(normalized.strip(), fmt).replace(tzinfo=SHANGHAI)
        except ValueError:
            continue
    try:
        parsed = datetime.fromisoformat(normalized)
        return parsed.replace(tzinfo=parsed.tzinfo or SHANGHAI)
    except ValueError:
        return None


def _side(value: object) -> str | None:
    text = str(value or "").strip().casefold()
    if any(token in text for token in ("买入", "证券买入", "buy")):
        return "BUY"
    if any(token in text for token in ("卖出", "证券卖出", "sell")):
        return "SELL"
    return None


def _asset_type(value: object) -> str | None:
    text = str(value or "").strip().casefold()
    if not text:
        return None
    if "etf" in text:
        return "ETF"
    if "基金" in text:
        return "MUTUAL_FUND"
    if "转债" in text or "可转" in text:
        return "CONVERTIBLE_BOND"
    if "股票" in text or "stock" in text:
        return "STOCK"
    return "OTHER"


def _preview_row(row_number: int, row: dict[str, object], mapping: dict[str, str], confidence: Decimal = Decimal("1")) -> TradePreviewRow:
    def get(field: str) -> object:
        return row.get(mapping.get(field, ""), "")
    traded_at = _datetime(get("traded_at"))
    quantity = _decimal(get("quantity"))
    price = _decimal(get("price_cny"))
    reported_amount = _decimal(get("gross_amount_cny"))
    amount = reported_amount if reported_amount is not None else (quantity * price if quantity and price else None)
    currency = str(get("currency") or "CNY").strip().upper()
    proposed: dict[str, Any] = {
        "traded_at": traded_at.isoformat() if traded_at else None,
        "security_code": str(get("security_code") or "").strip() or None,
        "security_name": str(get("security_name") or "").strip() or None,
        "side": _side(get("side")),
        "quantity": str(quantity) if quantity is not None else None,
        "price_cny": str(price) if price is not None else None,
        "gross_amount_cny": str(amount) if amount is not None else None,
        "fee_cny": str(_decimal(get("fee_cny")) or Decimal("0")),
        "asset_type": _asset_type(get("asset_type")),
        "currency": currency,
        "account_alias": str(get("account_alias") or "").strip() or "默认账户",
        "broker_trade_id": str(get("broker_trade_id") or "").strip() or None,
        "account_value_cny": str(_decimal(get("account_value_cny"))) if _decimal(get("account_value_cny")) is not None else None,
        "source_row": row_number,
        "source_confidence": str(confidence),
    }
    issues: list[str] = []
    if currency not in {"CNY", "人民币", "RMB"}:
        issues.append("仅支持人民币 CNY 交易")
    if traded_at is None:
        issues.append("缺少或无法识别交易时间")
    if not proposed["security_code"] and not proposed["security_name"]:
        issues.append("缺少证券代码或名称")
    if proposed["side"] is None:
        issues.append("无法识别买卖方向")
    if quantity is None or quantity <= 0:
        issues.append("成交数量必须大于 0")
    if price is None or price <= 0:
        issues.append("成交价格必须大于 0")
    if quantity and price and reported_amount is not None:
        expected = quantity * price
        if abs(reported_amount - expected) > max(Decimal("0.01"), expected * Decimal("0.001")):
            issues.append("成交金额与数量乘价格偏差超过 0.1%")
    if layout_issue := str(row.get("版式校验") or "").strip():
        issues.append(layout_issue)
    if confidence < OCR_CONFIDENCE_THRESHOLD:
        issues.append("OCR 置信度低于 85%，需要人工复核")
    overbound = currency not in {"CNY", "人民币", "RMB"}
    return TradePreviewRow(
        row_number=row_number,
        raw_values={str(key): None if value is None else str(value) for key, value in row.items()},
        proposed=proposed,
        confidence=confidence,
        status="OVERBOUND" if overbound else "REVIEW_REQUIRED" if issues else "PASS",
        issues=tuple(issues),
    )


def _decode_csv(data: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ImportParseError("CSV encoding must be UTF-8, UTF-8 BOM or GB18030")


def _table_preview(filename: str, data: bytes, selected_sheet: str | None) -> tuple[str, list[str], list[dict[str, object]], tuple[str, ...], str | None]:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if len(data) > MAX_TABLE_BYTES:
        raise ImportLimitError("table file exceeds 10 MiB")
    if suffix == "csv":
        text = _decode_csv(data)
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        columns = [str(item or "").strip() for item in (reader.fieldnames or [])]
        rows = [{str(key or "").strip(): value for key, value in row.items()} for row in reader]
        return "CSV", columns, rows, (), None
    if suffix != "xlsx":
        raise ImportParseError("only CSV and XLSX table files are supported")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportParseError("XLSX support is not installed") from exc
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportParseError("XLSX workbook could not be opened") from exc
    sheets = tuple(workbook.sheetnames)
    if not sheets:
        raise ImportParseError("XLSX workbook has no sheets")
    sheet_name = selected_sheet if selected_sheet in sheets else sheets[0]
    sheet = workbook[sheet_name]
    values = sheet.iter_rows(values_only=True)
    try:
        header = next(values)
    except StopIteration as exc:
        raise ImportParseError("XLSX sheet is empty") from exc
    columns = [str(item or "").strip() for item in header]
    rows = [dict(zip(columns, row)) for row in values if any(value not in (None, "") for value in row)]
    workbook.close()
    return "XLSX", columns, rows, sheets, sheet_name


def _statement_rows(
    grouped: list[list[dict[str, object]]],
) -> list[tuple[dict[str, object], Decimal]] | None:
    """Parse deterministic two-line mobile statement rows such as 同花顺对账单."""
    header_index = next((index for index, row in enumerate(grouped) if all(
        token in _key("".join(str(cell["text"]) for cell in row))
        for token in ("本月操作", "价格数量", "金额税费")
    )), None)
    if header_index is None:
        return None

    header = grouped[header_index]
    middle_header = next(cell for cell in header if "价格数量" in _key(cell["text"]))
    right_header = next(cell for cell in header if "金额税费" in _key(cell["text"]))
    middle_x = float(middle_header["x"])
    right_x = float(right_header["x"])
    current_year_month: tuple[int, int] | None = None
    parsed: list[tuple[dict[str, object], Decimal]] = []

    def column_decimal(row: list[dict[str, object]], target_x: float) -> Decimal | None:
        numeric = [(cell, _decimal(cell["text"])) for cell in row]
        candidates = [(cell, value) for cell, value in numeric if value is not None]
        if not candidates:
            return None
        cell, value = min(candidates, key=lambda item: abs(float(item[0]["x"]) - target_x))
        other_x = right_x if target_x == middle_x else middle_x
        return value if abs(float(cell["x"]) - target_x) < abs(float(cell["x"]) - other_x) else None

    rows = grouped[header_index + 1:]
    index = 0
    while index < len(rows):
        row = rows[index]
        joined = " ".join(str(cell["text"]) for cell in row)
        month_match = re.search(r"(20\d{2})\s*[-/.年]\s*(\d{1,2})", joined)
        if month_match:
            year, month = int(month_match.group(1)), int(month_match.group(2))
            if 1 <= month <= 12:
                current_year_month = (year, month)
            index += 1
            continue

        operation_cell = next((cell for cell in row if re.search(
            r"(?:证券)?(?:买入|卖出)\s*[-—－一:：]", str(cell["text"])
        )), None)
        if operation_cell is None or current_year_month is None or index + 1 >= len(rows):
            index += 1
            continue
        operation_match = re.search(
            r"(?:证券)?(买入|卖出)\s*[-—－一:：]\s*(.+)", str(operation_cell["text"])
        )
        if operation_match is None:
            index += 1
            continue

        detail = rows[index + 1]
        detail_text = " ".join(str(cell["text"]) for cell in detail)
        date_match = re.search(
            r"(?:买|卖)?\s*(\d{1,2})\s*[-/.]\s*(\d{1,2})\s*(\d{1,2})\s*:\s*(\d{2})",
            detail_text,
        )
        price = column_decimal(row, middle_x)
        net_amount = column_decimal(row, right_x)
        quantity = column_decimal(detail, middle_x)
        fee = column_decimal(detail, right_x)
        if date_match is None or price is None or quantity is None:
            index += 1
            continue
        transaction_month, day, hour, minute = (int(value) for value in date_match.groups())
        year = current_year_month[0]
        try:
            traded_at = datetime(year, transaction_month, day, hour, minute, tzinfo=SHANGHAI)
        except ValueError:
            index += 1
            continue
        gross_amount = abs(quantity * price)
        side = "买入" if operation_match.group(1) == "买入" else "卖出"
        expected_net = gross_amount + (fee or Decimal("0")) if side == "买入" else gross_amount - (fee or Decimal("0"))
        statement_issues: list[str] = []
        if transaction_month != current_year_month[1]:
            statement_issues.append("交易月份与对账单月份分组不一致")
        if net_amount is not None:
            tolerance = max(Decimal("0.01"), gross_amount * Decimal("0.001"))
            if abs(abs(net_amount) - expected_net) > tolerance:
                statement_issues.append("截图净发生额与成交金额及税费不一致")
        confidence = min(Decimal(str(cell["score"])) for cell in (*row, *detail))
        parsed.append(({
            "成交时间": traded_at.strftime("%Y-%m-%d %H:%M:%S"),
            "证券名称": operation_match.group(2).strip(),
            "买卖方向": side,
            "成交数量": str(abs(quantity)),
            "成交价格": str(abs(price)),
            "成交金额": str(gross_amount),
            "手续费": str(abs(fee or Decimal("0"))),
            "资产类型": "股票",
            "币种": "CNY",
            "对账单净发生额": str(net_amount) if net_amount is not None else "",
            "版式校验": "；".join(statement_issues),
        }, confidence))
        index += 2
    return parsed


def _ocr_rows(data: bytes) -> list[tuple[dict[str, object], Decimal]]:
    from PIL import Image
    from app.llm.ocr_portfolio_parser import OCRPortfolioParser

    try:
        image = Image.open(io.BytesIO(data)).convert("RGB")
    except Exception as exc:
        raise ImportParseError("trade screenshot could not be opened") from exc
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    result, _ = OCRPortfolioParser.get_instance()._get_engine()(buffer.getvalue())
    if not result:
        return []
    boxes = []
    for box, text, score in result:
        boxes.append({
            "text": str(text).strip(),
            "score": Decimal(str(score)),
            "x": (box[0][0] + box[2][0]) / 2,
            "y": (box[0][1] + box[2][1]) / 2,
            "height": abs(box[2][1] - box[0][1]),
        })
    tolerance = float(median([item["height"] for item in boxes])) * 0.45
    grouped: list[list[dict[str, object]]] = []
    for item in sorted(boxes, key=lambda cell: (cell["y"], cell["x"])):
        if not grouped or abs(float(item["y"]) - sum(float(cell["y"]) for cell in grouped[-1]) / len(grouped[-1])) > tolerance:
            grouped.append([item])
        else:
            grouped[-1].append(item)
    grouped = [sorted(row, key=lambda cell: cell["x"]) for row in grouped]
    statement = _statement_rows(grouped)
    if statement is not None:
        return statement
    header_index = next((index for index, row in enumerate(grouped) if any(
        token in "".join(str(cell["text"]) for cell in row)
        for token in ("成交时间", "成交日期", "证券代码", "买卖方向")
    )), None)
    if header_index is None:
        return []
    headers = grouped[header_index]
    rows: list[tuple[dict[str, object], Decimal]] = []
    for row in grouped[header_index + 1:]:
        cells: dict[str, object] = {}
        scores: list[Decimal] = []
        for cell in row:
            header = min(headers, key=lambda candidate: abs(float(candidate["x"]) - float(cell["x"])))
            cells[str(header["text"])] = cell["text"]
            scores.append(cell["score"])
        if cells:
            rows.append((cells, min(scores) if scores else Decimal("0")))
    return rows


def preview_trade_files(files: list[tuple[str, str, bytes]], *, selected_sheet: str | None = None) -> TradeImportPreview:
    if not files:
        raise ImportParseError("at least one file is required")
    digest = sha256()
    for filename, content_type, data in files:
        digest.update(filename.encode("utf-8", errors="ignore"))
        digest.update(b"\x00")
        digest.update(data)
        digest.update(b"\x00")
    image_types = {"image/png", "image/jpeg", "image/webp"}
    are_images = all(content_type in image_types for _, content_type, _ in files)
    if are_images:
        if len(files) > MAX_IMAGES:
            raise ImportLimitError("image batch exceeds 10 files")
        if any(not data or len(data) > MAX_IMAGE_BYTES for _, _, data in files):
            raise ImportLimitError("each image must be between 1 byte and 5 MiB")
        extracted: list[tuple[dict[str, object], Decimal]] = []
        for _, _, data in files:
            extracted.extend(_ocr_rows(data))
        if not extracted:
            raise ImportParseError("未识别到可导入的交易明细；请上传包含展开交易行的完整截图")
        columns = list(dict.fromkeys(key for row, _ in extracted for key in row))
        mapping = _mapping(columns)
        preview_rows = tuple(_preview_row(index, row, mapping, confidence) for index, (row, confidence) in enumerate(extracted, 1))
        source_type, sheets, active_sheet = "IMAGE", (), None
    else:
        if len(files) != 1:
            raise ImportParseError("table imports accept exactly one CSV or XLSX file")
        filename, _, data = files[0]
        source_type, columns, rows, sheets, active_sheet = _table_preview(filename, data, selected_sheet)
        if len(rows) > MAX_ROWS:
            raise ImportLimitError("table exceeds 20,000 data rows")
        mapping = _mapping(columns)
        preview_rows = tuple(_preview_row(index, row, mapping) for index, row in enumerate(rows, 2))
    if len(preview_rows) > MAX_ROWS:
        raise ImportLimitError("preview exceeds 20,000 rows")
    return TradeImportPreview(
        source_type=source_type,
        source_digest=digest.hexdigest(),
        file_count=len(files),
        detected_columns=tuple(columns),
        suggested_mapping=mapping,
        sheets=sheets,
        selected_sheet=active_sheet,
        rows=preview_rows,
        accepted_count=sum(item.status == "PASS" for item in preview_rows),
        review_count=sum(item.status == "REVIEW_REQUIRED" for item in preview_rows),
        rejected_count=sum(item.status == "OVERBOUND" for item in preview_rows),
    )
