from argparse import ArgumentParser
from collections import Counter
from hashlib import sha256
from pathlib import Path
import json
import re

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree
from PIL import Image
import pdfplumber
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "output/s2c-authoring"


def normalized(text):
    return re.sub(r"\s+", "", text)


def main():
    parser = ArgumentParser()
    parser.add_argument("--template", type=Path, required=True)
    args = parser.parse_args()
    report = json.loads((WORK / "build-report.json").read_text(encoding="utf-8"))
    target = Path(report["output"])
    doc = Document(target)
    reference = Document(args.template)
    actual = []
    for element in doc._element.body:
        if element.tag == qn("w:p"):
            text = Paragraph(element, doc._body).text
            if text:
                actual.append(text)
        elif element.tag == qn("w:tbl"):
            for row in Table(element, doc._body).rows:
                actual.extend(cell.text for cell in row.cells)
    if actual != report["expected_text"]:
        raise ValueError("DOCX text does not match Markdown content")
    chapters = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    expected = ["3.1 项目总体设计", "3.2 界面设计", "3.3 数据库设计", "3.4 系统关键技术实现", "3.5 系统主要界面"]
    if chapters != expected:
        raise ValueError("chapter scope differs from request")
    fulltext = "\n".join(actual)
    forbidden = [term for term in ["创新", "3.6", "不是", "而是", "契约", "栈", "落", "死", "拆", "偏"] if term in fulltext]
    if forbidden:
        raise ValueError(f"unexpected document terms: {forbidden}")
    if len(doc.tables) != report["tables"] or len(doc.inline_shapes) != len(report["figures"]):
        raise ValueError("table or figure count differs from source")
    captions = [item.text for item in doc.paragraphs if item.text.startswith("图 3-")]
    numbers = [int(re.match(r"图 3-(\d+)", caption).group(1)) for caption in captions]
    if numbers != list(range(1, len(report["figures"]) + 1)):
        raise ValueError("figure numbering must be continuous")
    for key in ["page_width", "page_height", "top_margin", "bottom_margin", "left_margin", "right_margin"]:
        if getattr(doc.sections[0], key) != getattr(reference.sections[-1], key):
            raise ValueError(f"reference geometry mismatch: {key}")
    font_sizes = {"Title": 26, "Heading 1": 18, "Heading 2": 15, "Heading 3": 14}
    for item in doc.paragraphs:
        if item.style.name in font_sizes:
            if any(run.font.size.pt != font_sizes[item.style.name] for run in item.runs if run.text):
                raise ValueError("heading font size mismatch")
        if item.text and any(run.font.name != "楷体" for run in item.runs if run.text):
            raise ValueError("paragraph font family mismatch")
    printable = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    for table in doc.tables:
        if len(table.columns) not in {4, 5}:
            raise ValueError("unexpected column count")
        if abs(sum(column.width for column in table.columns) - printable) > 635 * len(table.columns):
            raise ValueError("table exceeds printable width")
        if table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is None:
            raise ValueError("table header must repeat")
        if any(row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is None for row in table.rows):
            raise ValueError("table row may split between pages")
    image_references = doc._element.xpath(".//a:blip/@r:embed")
    if len(image_references) != len(report["figures"]) or any(key not in doc.part.rels for key in image_references):
        raise ValueError("figure relationship missing")
    for figure, key in zip(report["figures"], image_references, strict=True):
        relationship = doc.part.rels[key]
        if sha256(relationship.target_part.blob).hexdigest() != figure["sha256"]:
            raise ValueError("figure bytes changed")
        if Path(figure["source"]).parent.name == "s2c":
            with Image.open(figure["source"]) as image:
                printed_font = 40 * figure["width_mm"] / image.width * 72 / 25.4
            if printed_font < 7.5:
                raise ValueError("diagram node labels are too small in Word")
    pdf_path = WORK / (target.stem + ".pdf")
    pdf = PdfReader(pdf_path)
    pages = []
    pdf_text = []
    with pdfplumber.open(pdf_path) as rendered:
        for index, page in enumerate(rendered.pages, 1):
            text = page.extract_text(x_tolerance=1, y_tolerance=3) or ""
            pdf_text.append(text)
            outside = [character for character in page.chars
                       if character["x0"] < -1 or character["x1"] > page.width + 1
                       or character["top"] < -1 or character["bottom"] > page.height + 1]
            outside_images = [item for item in page.images
                              if item["x0"] < -1 or item["x1"] > page.width + 1
                              or item["top"] < -1 or item["bottom"] > page.height + 1]
            if outside or outside_images:
                raise ValueError(f"content exceeds page boundary: {index}")
            for image in page.images:
                overlapping = [character for character in page.chars
                               if character["x0"] < image["x1"] - 0.5
                               and character["x1"] > image["x0"] + 0.5
                               and character["top"] < image["bottom"] - 0.5
                               and character["bottom"] > image["top"] + 0.5]
                if overlapping:
                    raise ValueError(f"text overlaps figure boundary: {index}")
            if not text.strip() and not page.images:
                raise ValueError(f"unexpected blank page: {index}")
            pages.append({"page": index, "characters": len(page.chars), "images": len(page.images),
                          "first_text": text[:100], "page_boundary": "PASS"})
    joined = normalized("\n".join(pdf_text))
    if sum(page["images"] for page in pages) != len(report["figures"]):
        raise ValueError("PDF figure count differs from DOCX")
    for heading in report["headings"]:
        if normalized(heading) not in joined:
            raise ValueError(f"heading missing from PDF: {heading}")
    required = Counter(normalized(fulltext))
    available = Counter(joined)
    missing = {key: value - available[key] for key, value in required.items() if value > available[key]}
    if missing:
        raise ValueError(f"text characters missing from PDF: {missing}")
    if sha256(args.template.read_bytes()).hexdigest() != report["template_sha256"]:
        raise ValueError("reference changed")
    # SVG 文字通过 XML 库读取，验证复用图形的来源结构。
    svg_titles = {}
    provenance = json.loads((ROOT / "docs/submission/figures/s2c/figure-sources.json").read_text(encoding="utf-8"))
    for name, sources in provenance["sources"].items():
        if any(not (ROOT / source).is_file() for source in sources):
            raise ValueError("diagram code source is missing")
        if not (ROOT / "docs/submission/figures/s2c" / (name + ".mmd")).is_file():
            raise ValueError("Mermaid source is missing")
    for figure in report["figures"]:
        vector = Path(figure["source"]).with_suffix(".svg")
        if vector.is_file():
            tree = etree.parse(str(vector))
            svg_titles[str(vector.relative_to(ROOT))] = tree.xpath("//*[local-name()='text']//text()")
    original_architecture = etree.parse(str(ROOT / "docs/submission/Prism系统技术架构.svg"))
    document_architecture = etree.parse(str(ROOT / "docs/submission/figures/s2c-shared/system-architecture.svg"))
    if original_architecture.xpath("//*[local-name()='text']//text()") != document_architecture.xpath("//*[local-name()='text']//text()"):
        raise ValueError("architecture figure text changed")
    geometry = json.loads((WORK / "diagram-overlap-audit.json").read_text(encoding="utf-8"))
    diagrams = [figure for figure in report["figures"] if Path(figure["source"]).parent.name != "s2c-ui"]
    if geometry["status"] != "PASS" or len(geometry["figures"]) != len(diagrams):
        raise ValueError("diagram geometry checks did not pass")
    expected_vectors = {Path(figure["source"]).with_suffix(".svg").resolve() for figure in diagrams}
    checked_vectors = {(ROOT / figure["source"]).resolve() for figure in geometry["figures"]}
    if expected_vectors != checked_vectors:
        raise ValueError("geometry checks must cover every document diagram")
    for figure in geometry["figures"]:
        vector = ROOT / figure["source"]
        if sha256(vector.read_bytes()).hexdigest() != figure["svg_sha256"] or sha256(vector.with_suffix(".png").read_bytes()).hexdigest() != figure["png_sha256"]:
            raise ValueError("diagram changed after geometry checks")
    capture = json.loads((ROOT / "docs/submission/figures/s2c-ui/capture-manifest.json").read_text(encoding="utf-8"))
    if capture["snapshot_adapter"] or capture["browser_errors"]:
        raise ValueError("current interface capture did not pass")
    for screenshot in capture["screenshots"]:
        asset = ROOT / "docs/submission/figures/s2c-ui" / screenshot["filename"]
        if sha256(asset.read_bytes()).hexdigest() != screenshot["sha256"]:
            raise ValueError("interface screenshot differs from capture")
    for filename, expected_hash in capture["static_sha256"].items():
        if sha256((ROOT / "app/api/static" / filename).read_bytes()).hexdigest() != expected_hash:
            raise ValueError("interface code changed after capture")
    result = {"status": "PASS", "document": str(target), "pages": len(pdf.pages),
              "chapters": chapters, "tables": len(doc.tables), "figures": len(doc.inline_shapes),
              "text_matches_source": True, "source_characters": len(normalized(fulltext)),
              "reference_unchanged": True, "template_geometry": "PASS", "heading_fonts": "PASS",
              "pdf_text_coverage": "PASS", "page_boundaries": "PASS", "page_details": pages,
              "svg_text": svg_titles, "current_interface_capture": "PASS", "diagram_geometry": "PASS",
              "visual_review": "NOT_PERFORMED_USER_RESTRICTION",
              "sha256": sha256(target.read_bytes()).hexdigest()}
    (WORK / "verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({key: result[key] for key in ["status", "pages", "tables", "figures", "source_characters",
                                                  "reference_unchanged", "pdf_text_coverage", "page_boundaries"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
