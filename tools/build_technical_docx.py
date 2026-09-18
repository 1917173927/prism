from pathlib import Path
import argparse
import re
from copy import deepcopy

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from markdown_it import MarkdownIt
from latex2mathml.converter import convert
from lxml import etree
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/submission/competition-technical-solution.md'
TARGET = ROOT / 'docs/submission/Prism-技术文档.docx'


def font(run, size=12, bold=False):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '楷体')


def field(paragraph, instruction):
    run = paragraph.add_run()
    node = OxmlElement('w:fldSimple')
    node.set(qn('w:instr'), instruction)
    run._r.addnext(node)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--template', required=True)
    parser.add_argument('--output', type=Path, default=TARGET)
    args = parser.parse_args()
    doc = Document(args.template)
    table_properties = deepcopy(doc.tables[0]._tbl.tblPr)
    body = doc._element.body
    section = deepcopy(body.sectPr)
    for child in list(body):
        body.remove(child)
    body.append(section)
    for relationship_id, relationship in list(doc.part.rels.items()):
        if relationship.reltype.endswith(('/image', '/header', '/footer')):
            doc.part.drop_rel(relationship_id)
    for tag in ('w:headerReference', 'w:footerReference'):
        for element in list(section.findall(qn(tag))):
            section.remove(element)
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2)
        sec.header_distance = sec.footer_distance = Cm(1)
        sec.different_first_page_header_footer = True
        header = sec.header.paragraphs[0]
        font(header.add_run('Prism 个性化证券投顾智能体系统'), 9)
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        field(footer, 'PAGE')
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '楷体')
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(0.85)
    for level, size in [(1,20),(2,16),(3,14),(4,12)]:
        if f'Heading {level}' not in doc.styles:
            doc.styles.add_style(f'Heading {level}', WD_STYLE_TYPE.PARAGRAPH)
        style = doc.styles[f'Heading {level}']
        style.font.size = Pt(size)
        style.font.name = 'Times New Roman'
        style.font.bold = True
        style.font.color.rgb = RGBColor(0,0,0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '楷体')
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.page_break_before = level == 1
        outline = OxmlElement('w:outlineLvl')
        outline.set(qn('w:val'), str(level-1))
        style._element.get_or_add_pPr().append(outline)
    if 'Title' not in doc.styles:
        doc.styles.add_style('Title', WD_STYLE_TYPE.PARAGRAPH)
    for level in range(1,5):
        name = next((s.name for s in doc.styles if s.name.lower() == f'toc {level}'), f'toc {level}')
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style = doc.styles[name]
        style.font.size = Pt(10.5)
        style.font.name = 'Times New Roman'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '楷体')
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.keep_with_next = False
    md = MarkdownIt('commonmark').enable('table')
    tokens = md.parse(SOURCE.read_text(encoding='utf-8'))
    transform = etree.XSLT(etree.parse('C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL'))
    index = 0
    bullet = False
    table = None
    row = None
    col = 0
    target_paragraph = None
    references = False
    appendix = False
    ordered_number = None
    while index < len(tokens):
        token = tokens[index]
        if token.type == 'heading_open':
            value = tokens[index+1].content
            level = int(token.tag[1:]) - 1
            if level == 0:
                p = doc.add_paragraph(style='Title')
                p.paragraph_format.space_before = Pt(100)
                p.paragraph_format.first_line_indent = Cm(0)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                font(p.add_run('Prism\n个性化证券投顾智能体系统\n技术文档'), 24, True)
                doc.add_page_break()
                p = doc.add_paragraph('目录')
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                font(p.runs[0],20,True)
                field(doc.add_paragraph(), 'TOC \\o "1-4" \\h \\z \\u')
            else:
                p = doc.add_paragraph(style=f'Heading {level}')
                font(p.add_run(value), {1:20,2:16,3:14,4:12}[level],True)
                references = value == '第八章 参考文献'
                if level == 1:
                    appendix = value.startswith('附录')
            index += 3
            continue
        if token.type == 'bullet_list_open':
            bullet = True
        elif token.type == 'bullet_list_close':
            bullet = False
        elif token.type == 'ordered_list_open':
            ordered_number = int(token.attrGet('start') or 1)
        elif token.type == 'ordered_list_close':
            ordered_number = None
        elif token.type == 'table_open':
            table = doc.add_table(rows=0, cols=0)
            table._tbl.remove(table._tbl.tblPr)
            table._tbl.insert(0, deepcopy(table_properties))
            table.autofit = False
        elif token.type == 'tr_open':
            if len(table.columns) == 0:
                count = 0
                for item in tokens[index+1:]:
                    if item.type == 'tr_close':
                        break
                    count += item.type in ('th_open','td_open')
                widths = [2.5,7,4.5,3] if appendix else [17/count]*count
                for width in widths:
                    table.add_column(Cm(width))
            row = table.add_row()
            col = 0
            setting = OxmlElement('w:cantSplit')
            row._tr.get_or_add_trPr().append(setting)
            if len(table.rows) == 1:
                row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        elif token.type in ('th_open','td_open'):
            target_paragraph = row.cells[col].paragraphs[0]
            target_paragraph.paragraph_format.first_line_indent = Cm(0)
            target_paragraph.paragraph_format.space_after = Pt(3)
            target_paragraph.paragraph_format.space_before = Pt(3)
            target_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            col += 1
        elif token.type == 'table_close':
            if len(table.rows) <= 7:
                for table_row in table.rows:
                    for cell in table_row.cells:
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.keep_with_next = True
            table = None
            target_paragraph = None
        elif token.type == 'paragraph_open':
            if table is None:
                target_paragraph = doc.add_paragraph()
                if bullet:
                    target_paragraph.paragraph_format.first_line_indent = Cm(0)
                    target_paragraph.paragraph_format.left_indent = Cm(0.6)
                    font(target_paragraph.add_run('• '))
                elif ordered_number is not None:
                    font(target_paragraph.add_run(f'{ordered_number}. '))
                    ordered_number += 1
        elif token.type == 'inline':
            p = target_paragraph
            if p is None:
                raise ValueError(token.content)
            value = token.content
            if value.startswith('$$') and value.endswith('$$'):
                equation = value[2:-2]
                numbered = re.fullmatch(r'(.*)\\qquad \((\d+-\d+)\)', equation)
                if numbered:
                    equation = numbered.group(1)
                    p.paragraph_format.tab_stops.add_tab_stop(Cm(8.5), WD_TAB_ALIGNMENT.CENTER)
                    p.paragraph_format.tab_stops.add_tab_stop(Cm(17), WD_TAB_ALIGNMENT.RIGHT)
                    p.add_run('\t')
                math = etree.fromstring(convert(equation).encode())
                p._p.append(transform(math).getroot())
                if numbered:
                    font(p.add_run(f'\t({numbered.group(2)})'))
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if numbered else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.keep_with_next = True
            elif token.children and token.children[0].type == 'image':
                file = SOURCE.parent / token.children[0].attrGet('src')
                with Image.open(file) as img:
                    width, height = img.size
                max_height = 21 if file.stem in {'frontend-framework', 'backend-framework', 'context-dataflow'} else (14 if 'judge-06' in file.name else (16 if 'judge-' in file.name else 17.5))
                draw_width = min(17, max_height*width/height)
                p.paragraph_format.first_line_indent = Cm(0)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(str(file),width=Cm(draw_width))
            else:
                size = 10 if table is not None else 12
                if re.match(r'^[图表] [\dA]+-', value):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.first_line_indent = Cm(0)
                    size = 10.5
                if references:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.first_line_indent = Cm(-0.7)
                    p.paragraph_format.left_indent = Cm(0.7)
                font(p.add_run(value),size,table is not None and len(table.rows)==1)
        index += 1
    for p in doc.paragraphs:
        p.paragraph_format.widow_control = True
    doc.core_properties.title = 'Prism 个性化证券投顾智能体系统技术文档'
    doc.core_properties.author = 'Prism'
    doc.core_properties.subject = '系统设计 核心技术 创新与验证'
    doc.save(args.output)
    print(args.output)


if __name__ == '__main__':
    main()
