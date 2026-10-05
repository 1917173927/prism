from argparse import ArgumentParser
from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from urllib.parse import unquote
import json

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from markdown_it import MarkdownIt
from lxml import etree
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/submission/Prism-S2C-技术路线及实现方案.md"
TARGET = SOURCE.with_name("Prism-S2C-技术路线及实现方案-图示版.docx")
REFERENCE_HASH = "0ad3cd9fce58ce2b31562b5bc9412d03d8a6cd5d0ab97c63ccddf8aba8c7f228"
REPORT = ROOT / "output/s2c-authoring/build-report.json"


def set_font(run, size, bold=False):
    run.font.name = "楷体"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "楷体")


def inline_text(token):
    return "".join(child.content if child.type in {"text", "code_inline"} else "\n"
                   if child.type in {"softbreak", "hardbreak"} else ""
                   for child in token.children or [])


def add_inline(paragraph, token, *, size=12, bold=False):
    for child in token.children or []:
        if child.type in {"text", "code_inline"}:
            set_font(paragraph.add_run(child.content), size, bold)
        elif child.type in {"softbreak", "hardbreak"}:
            paragraph.add_run().add_break()
        elif child.type in {"strong_open", "strong_close", "em_open", "em_close", "link_open", "link_close"}:
            continue
        else:
            raise ValueError(f"unsupported inline token: {child.type}")


def set_table_width(table, widths):
    properties = table._tbl.tblPr
    width = properties.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        properties.append(width)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(round(sum(widths) * 1440 / 25.4)))
    indent = properties.find(qn("w:tblInd"))
    if indent is not None:
        properties.remove(indent)
    table.autofit = False
    for column, size in zip(table.columns, widths, strict=True):
        column.width = Mm(size)
    for row in table.rows:
        for cell, size in zip(row.cells, widths, strict=True):
            cell.width = Mm(size)


def main():
    parser = ArgumentParser()
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=TARGET)
    args = parser.parse_args()
    original = args.template.read_bytes()
    if sha256(original).hexdigest() != REFERENCE_HASH:
        raise ValueError("reference document changed")
    doc = Document(args.template)
    reference = {"title": deepcopy(doc.paragraphs[187]._p),
                 "chapter": deepcopy(doc.paragraphs[188]._p),
                 "section": deepcopy(doc.paragraphs[189]._p),
                 "minor": deepcopy(doc.paragraphs[200]._p),
                 "body": deepcopy(doc.paragraphs[201]._p),
                 "caption": deepcopy(doc.paragraphs[192]._p)}
    table_properties = deepcopy(doc.tables[0]._tbl.tblPr)
    body = doc._element.body
    section = deepcopy(body.sectPr)
    for item in list(body):
        body.remove(item)
    body.append(section)
    # 新正文复用参考文件格式，移除旧正文的图片关系。
    for identifier, relationship in list(doc.part.rels.items()):
        if relationship.reltype.endswith("/image"):
            doc.part.drop_rel(identifier)

    sizes = {"title": 26, "chapter": 18, "section": 15, "minor": 14, "body": 12, "caption": 10.5}
    styles = {"title": "Title", "chapter": "Heading 1", "section": "Heading 2", "minor": "Heading 3"}
    for role, style_name in styles.items():
        if style_name not in doc.styles:
            doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style = doc.styles[style_name]
        style.font.name = "楷体"
        style.font.size = Pt(sizes[role])
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "楷体")
        style.paragraph_format.page_break_before = False
        style.paragraph_format.keep_with_next = True
        borders = style._element.get_or_add_pPr().find(qn("w:pBdr"))
        if borders is not None:
            style._element.get_or_add_pPr().remove(borders)

    def paragraph(role):
        node = deepcopy(reference[role])
        body.insert(len(body) - 1, node)
        from docx.text.paragraph import Paragraph
        result = Paragraph(node, doc._body)
        result.clear()
        result.style = styles.get(role, "Normal")
        result.paragraph_format.widow_control = True
        result.paragraph_format.keep_with_next = role != "body"
        result.paragraph_format.keep_together = role != "body"
        result.paragraph_format.page_break_before = False
        if role in styles:
            result.paragraph_format.first_line_indent = Pt(0)
            result.paragraph_format.left_indent = Pt(0)
            result.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if role == "caption":
            result.alignment = WD_ALIGN_PARAGRAPH.CENTER
            result.paragraph_format.first_line_indent = Pt(0)
            result.paragraph_format.left_indent = Pt(0)
            result.paragraph_format.keep_with_next = False
        return result

    tokens = MarkdownIt("commonmark").enable("table").parse(SOURCE.read_text(encoding="utf-8"))
    headings, figures, expected = [], [], []
    index = 0
    while index < len(tokens):
        token = tokens[index]
        if token.type == "heading_open":
            text = inline_text(tokens[index + 1])
            role = {"h1": "title", "h2": "chapter", "h3": "section", "h4": "minor"}[token.tag]
            item = paragraph(role)
            set_font(item.add_run(text), sizes[role], True)
            headings.append(text)
            expected.append(text)
            index += 3
            continue
        if token.type == "paragraph_open":
            content = tokens[index + 1]
            images = [child for child in content.children or [] if child.type == "image"]
            if images:
                if len(images) != 1 or len(content.children) != 1:
                    raise ValueError("image must have its own paragraph")
                asset = (SOURCE.parent / unquote(images[0].attrGet("src"))).resolve()
                with Image.open(asset) as image:
                    image.verify()
                with Image.open(asset) as image:
                    pixel_width, pixel_height = image.size
                crop_top = 0.0
                if asset.stem in {"frontend-framework", "backend-framework", "context-dataflow"}:
                    vector = etree.parse(str(asset.with_suffix(".svg")))
                    crop_top = 32 / float(vector.getroot().get("height"))
                elif asset.stem in {"judge-04-evidence-validation", "judge-05-rebalancing", "judge-06-decision-gates"}:
                    crop_top = 0.07
                geometry = doc.sections[0]
                printable = geometry.page_width.mm - geometry.left_margin.mm - geometry.right_margin.mm
                visible_height = pixel_height * (1 - crop_top)
                width = min(printable, 210 * pixel_width / visible_height)
                item = paragraph("body")
                item.paragraph_format.first_line_indent = Pt(0)
                item.alignment = WD_ALIGN_PARAGRAPH.CENTER
                item.paragraph_format.keep_with_next = True
                item.paragraph_format.keep_together = True
                item.paragraph_format.space_after = Pt(20)
                shape = item.add_run().add_picture(str(asset), width=Mm(width), height=Mm(width * visible_height / pixel_width))
                if crop_top:
                    # 文档使用原图片字节，通过图片格式设置隐藏原图的页眉编号。
                    fill = shape._inline.graphic.graphicData.pic.blipFill
                    crop = OxmlElement("a:srcRect")
                    crop.set("t", str(round(crop_top * 100000)))
                    fill.insert(1, crop)
                figures.append({"source": str(asset), "sha256": sha256(asset.read_bytes()).hexdigest(),
                                "width_mm": width, "height_mm": width * visible_height / pixel_width,
                                "crop_top": crop_top})
            else:
                text = inline_text(content)
                caption = text.startswith(("图 3-", "表 3-"))
                item = paragraph("caption" if caption else "body")
                add_inline(item, content, size=10.5 if caption else 12)
                expected.append(text)
            index += 3
            continue
        if token.type == "table_open":
            rows, row = [], None
            index += 1
            while tokens[index].type != "table_close":
                current = tokens[index]
                if current.type == "tr_open":
                    row = []
                elif current.type == "inline":
                    if row is None:
                        raise ValueError("table row missing")
                    row.append(current)
                elif current.type == "tr_close":
                    rows.append(row)
                    row = None
                index += 1
            count = len(rows[0])
            if count not in {4, 5} or any(len(row) != count for row in rows):
                raise ValueError("tables require four or five columns")
            table = doc.add_table(rows=len(rows), cols=count)
            table._tbl.remove(table._tbl.tblPr)
            table._tbl.insert(0, deepcopy(table_properties))
            geometry = doc.sections[0]
            printable = geometry.page_width.mm - geometry.left_margin.mm - geometry.right_margin.mm
            first = 37 if any("数据表" in inline_text(cell) for cell in rows[0]) else 28
            widths = [first] + [(printable - first) / (count - 1)] * (count - 1)
            set_table_width(table, widths)
            for row_index, (table_row, values) in enumerate(zip(table.rows, rows, strict=True)):
                properties = table_row._tr.get_or_add_trPr()
                properties.append(OxmlElement("w:cantSplit"))
                if row_index == 0:
                    properties.append(OxmlElement("w:tblHeader"))
                for cell, value in zip(table_row.cells, values, strict=True):
                    item = cell.paragraphs[0]
                    item.paragraph_format.first_line_indent = Pt(0)
                    item.paragraph_format.left_indent = Pt(0)
                    item.paragraph_format.space_before = Pt(3)
                    item.paragraph_format.space_after = Pt(3)
                    item.paragraph_format.line_spacing = 1.1
                    item.paragraph_format.keep_with_next = row_index == 0
                    item.paragraph_format.widow_control = True
                    add_inline(item, value, size=9.5, bold=row_index == 0)
                    expected.append(inline_text(value))
            index += 1
            continue
        raise ValueError(f"unsupported block token: {token.type}")

    doc.core_properties.title = "Prism S2C 技术路线及实现方案"
    doc.core_properties.subject = "项目总体设计、界面、数据库、关键技术与主要界面"
    doc.core_properties.author = "Prism"
    doc.core_properties.keywords = "Prism,S2C,技术路线,实现方案"
    doc.core_properties.comments = ""
    doc.save(args.output)
    if sha256(args.template.read_bytes()).hexdigest() != REFERENCE_HASH:
        raise ValueError("reference changed during build")
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(json.dumps({"template_sha256": REFERENCE_HASH, "source": str(SOURCE),
                                 "output": str(args.output), "headings": headings, "figures": figures,
                                 "expected_text": expected, "tables": len(doc.tables),
                                 "visual_review": "NOT_PERFORMED_USER_RESTRICTION"}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(args.output), "headings": len(headings), "tables": len(doc.tables),
                      "figures": len(figures)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
