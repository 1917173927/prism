import argparse
from pathlib import Path
import subprocess
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from markdown_it import MarkdownIt
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/submission/system-test-report.md"
TARGET = ROOT / "docs/submission/Prism-系统测试报告.docx"
FIGURES = ROOT / "docs/submission/test-evidence/figures"
WORK = ROOT / "output/system-test-report"


def figures(selected=None):
    sources = sorted(FIGURES.glob("*.mmd"))
    if selected:
        sources = [source for source in sources if source.stem == selected]
        assert len(sources) == 1, selected
    for source in sources:
        carrier = WORK / (source.stem + ".docx")
        Document().save(carrier)
        subprocess.run(["officecli", "add", str(carrier), "/body", "--type", "diagram",
                        "--prop", f"src={source}", "--prop", "render=image", "--prop", "width=16cm",
                        "--prop", "background=white"], check=True)
        subprocess.run(["officecli", "close", str(carrier)], check=True)
        with zipfile.ZipFile(carrier) as package:
            images = [name for name in package.namelist() if name.startswith(("word/media/", "media/")) and name.endswith(".png")]
            assert len(images) == 1, images
            (FIGURES / (source.stem + ".png")).write_bytes(package.read(images[0]))


def set_font(run, size=11, bold=False, name="宋体"):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def field(paragraph, instruction):
    element = OxmlElement("w:fldSimple")
    element.set(qn("w:instr"), instruction)
    paragraph._p.append(element)


def style_table(table):
    table.autofit = False
    count = len(table.columns)
    widths = [1.12, 1.87, 1.80, 1.91] if count == 4 else [6.7 / count] * count
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement("w:" + side)
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            edge.set(qn("w:" + key), value)
        borders.append(edge)
    properties.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("top", "75"), ("bottom", "75"), ("left", "85"), ("right", "85")):
        node = OxmlElement("w:" + side)
        node.set(qn("w:w"), value)
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    properties.append(margins)
    for index, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if index == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for col, cell in enumerate(row.cells):
            cell.width = Inches(widths[col])
            cell.vertical_alignment = 1
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E7E9ED" if index == 0 else "FFFFFF" if index % 2 else "F7F8FA")
            cell._tc.get_or_add_tcPr().append(shade)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.1
                paragraph.paragraph_format.keep_with_next = index == 0
                for run in paragraph.runs:
                    set_font(run, 10.5, index == 0)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.72)
    sec.left_margin = sec.right_margin = Inches(.9)
    sec.header_distance = sec.footer_distance = Inches(.32)
    sec.different_first_page_header_footer = True
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(7)
    for name, size in (("Title", 25), ("Heading 1", 18), ("Heading 2", 14)):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "黑体")
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(9)
        style.paragraph_format.keep_with_next = True
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("Prism 系统测试报告"), 9)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(footer, "PAGE")

    tokens = MarkdownIt("commonmark").enable("table").parse(SOURCE.read_text(encoding="utf-8"))
    first_section = True
    pending_section = False
    current_table = None
    current_row = None
    cell_index = 0
    paragraph = None
    in_cell = False
    i = 0
    while i < len(tokens):
        token = tokens[i]
        if token.type == "heading_open":
            level = int(token.tag[1:])
            title = tokens[i+1].content
            if level == 1:
                paragraph = doc.add_paragraph(title, "Title")
                paragraph.paragraph_format.space_before = Pt(85)
                paragraph.paragraph_format.space_after = Pt(45)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                if level == 2 and first_section:
                    doc.add_page_break()
                    contents = doc.add_paragraph("目录")
                    set_font(contents.runs[0], 18, True, "黑体")
                    field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
                    first_section = False
                paragraph = doc.add_paragraph(title, f"Heading {level-1}")
                paragraph.paragraph_format.page_break_before = level == 2 or not pending_section
                pending_section = level == 2
            i += 3
            continue
        if token.type == "table_open":
            columns = 0
            cursor = i + 1
            while tokens[cursor].type != "tr_close":
                columns += tokens[cursor].type == "th_open"
                cursor += 1
            current_table = doc.add_table(rows=0, cols=columns)
        elif token.type == "tr_open":
            current_row = current_table.add_row()
            cell_index = 0
        elif token.type in ("th_open", "td_open"):
            paragraph = current_row.cells[cell_index].paragraphs[0]
            in_cell = True
            cell_index += 1
        elif token.type in ("th_close", "td_close"):
            in_cell = False
        elif token.type == "table_close":
            style_table(current_table)
            current_table = None
        elif token.type == "paragraph_open" and not in_cell:
            paragraph = doc.add_paragraph()
        elif token.type == "inline":
            for child in token.children or []:
                if child.type == "image":
                    path = SOURCE.parent / child.attrGet("src")
                    with Image.open(path) as picture:
                        width, height = picture.size
                    inches = min(6.5, 2.6 * width / height)
                    shape = paragraph.add_run().add_picture(str(path), width=Inches(inches))
                    shape._inline.docPr.set("descr", child.content)
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.keep_with_next = True
                    caption = doc.add_paragraph(child.content)
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    caption.paragraph_format.space_after = Pt(6)
                    for run in caption.runs:
                        set_font(run, 9.5)
                elif child.type in ("text", "code_inline"):
                    set_font(paragraph.add_run(child.content))
                elif child.type in ("softbreak", "hardbreak"):
                    paragraph.add_run(" ")
        elif token.type == "fence":
            for line in token.content.strip().splitlines():
                if line:
                    paragraph = doc.add_paragraph()
                    paragraph.paragraph_format.line_spacing = 1.1
                    set_font(paragraph.add_run(line), 10)
        i += 1
    doc.core_properties.title = "Prism 个性化证券投顾智能体系统测试报告"
    doc.core_properties.author = "Prism 项目组"
    doc.core_properties.subject = "A18 赛题系统测试与工程验证"
    doc.save(TARGET)
    print(str(TARGET))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--figures", action="store_true")
    parser.add_argument("--figure")
    args = parser.parse_args()
    WORK.mkdir(parents=True, exist_ok=True)
    if args.figures or args.figure:
        figures(args.figure)
    build()
