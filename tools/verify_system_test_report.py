from hashlib import sha256
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET
import zipfile

from docx import Document
import pymupdf
from markdown_it import MarkdownIt
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/submission/system-test-report.md"
DOCX = ROOT / "docs/submission/Prism-系统测试报告.docx"
EVIDENCE = ROOT / "docs/submission/test-evidence"
WORK = ROOT / "output/system-test-report"


def main():
    source = SOURCE.read_text(encoding="utf-8")
    results = json.loads((EVIDENCE / "current-results.json").read_text(encoding="utf-8"))
    http = json.loads((EVIDENCE / "http-results.json").read_text(encoding="utf-8"))
    suite = ET.parse(EVIDENCE / "profile-rules.xml").getroot().find("testsuite")
    assert suite.attrib["tests"] == "15"
    assert all(suite.attrib[name] == "0" for name in ("errors", "failures", "skipped"))
    assert len(results["tests"]) == 29 and all(row["status"] == "PASS" for row in results["tests"])
    assert len(http["tests"]) == 14 and all(row["status"] == "PASS" for row in http["tests"])
    assert results["source_sha256"] == sha256((ROOT / results["source_path"]).read_bytes()).hexdigest()
    assert http["performance"]["requests"] == http["performance"]["completed"] == 1000
    assert http["performance"]["sla_verified"] is False
    assert f'{http["performance"]["latency_ms"]["p95"]:,.3f}' in source
    assert f'{results["local_calculation_timing"]["p95_ms"]:.3f}' in source
    assert not re.search(r"栈|落|死|拆|偏|契约|不是.{0,40}而是|TODO|TBD|待填", source)
    doc = Document(DOCX)
    text = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert not re.search(r"TODO|TBD|待填|Error!|错误!|找不到目录", text)
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert len(headings) == 11, headings
    assert len(doc.inline_shapes) == 8
    assert all(shape._inline.docPr.get("descr") for shape in doc.inline_shapes)
    figures = []
    for token in MarkdownIt("commonmark").enable("table").parse(source):
        for child in token.children or []:
            if child.type == "image":
                image = SOURCE.parent / child.attrGet("src")
                with Image.open(image) as picture:
                    figures.append(dict(path=str(image.relative_to(ROOT)), width=picture.width, height=picture.height))
    assert len(figures) == 8
    with zipfile.ZipFile(DOCX) as package:
        xml = package.read("word/document.xml").decode("utf-8")
        assert "TOC" in xml
        assert all(table.rows[0]._tr.trPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader") is not None for table in doc.tables)
    pdf = pymupdf.open(WORK / "report.pdf")
    pages = []
    overflow = []
    for number, page in enumerate(pdf, 1):
        words = page.get_text("words")
        content = page.get_text()
        for word in words:
            if word[0] < 10 or word[1] < 8 or word[2] > page.rect.width-10 or word[3] > page.rect.height-8:
                overflow.append(dict(page=number, word=word[:5]))
        pages.append(dict(page=number, characters=len(content), beginning=content[:160], end=content[-120:]))
        assert len(content.strip()) > 200, f"正文不足的页面 {number}"
        page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5)).save(WORK / f"page-{number}.png")
    assert not overflow, overflow
    result = dict(status="PASS", verification_scope="证据、OOXML、Word 导出及 PDF 文字边界检查；未进行视觉识别",
                  pages=pages, page_count=len(pdf), figures=figures, tables=len(doc.tables), headings=headings,
                  docx_sha256=sha256(DOCX.read_bytes()).hexdigest(), overflow=overflow)
    (WORK / "document-qa.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(dict(status=result["status"], pages=len(pdf), figures=len(figures), tables=len(doc.tables)), ensure_ascii=False))


if __name__ == "__main__":
    main()
