"""Generate the formal Prism competition solution DOCX from the reviewed Markdown."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "submission" / "competition-technical-solution.md"
OUTPUT = ROOT / "docs" / "submission" / "Prism个性化证券投顾智能体系统项目详细方案.docx"
FIGURES = ROOT / "docs" / "submission" / "figures"

NAVY = "17365D"
PALE_BLUE = "EEF4F8"
LIGHT_BORDER = "D9D9D9"
BODY = "1F1F1F"
MUTED = "666666"

FIGURE_MAP = {
    "图 2-1": FIGURES / "prism-figure-01-overview.png",
    "图 3-1": FIGURES / "prism-figure-03-dependencies.png",
    "图 5-1": FIGURES / "prism-figure-02-architecture.png",
    "图 5-2": FIGURES / "prism-figure-04-research-dag.png",
    "图 5-4": FIGURES / "prism-figure-06-decision-gates.png",
}

FLOW_LABELS = {
    "图 5-3": ["原始记录", "证据", "来源校验", "事实", "研究发现", "双闸门", "建议回执"],
}

TOC_ENTRIES = [
    ("第一章  项目概况", 3),
    ("第二章  方案概要", 5),
    ("第三章  产品方案", 7),
    ("第四章  产品功能与应用展示", 9),
    ("第五章  技术方案", 14),
    ("第六章  创新亮点", 19),
    ("第七章  系统验证与项目成果", 21),
    ("第八章  落地模式与项目价值", 24),
    ("第九章  风险控制与发展规划", 26),
    ("附录  证据索引与实现说明", 28),
]


def set_run_font(run, east_asia="宋体", latin="Times New Roman", size=10.5, bold=None, color=BODY):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn("w:shd"))
    if old is not None:
        tc_pr.remove(old)
    tc_pr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill}"/>'))


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), LIGHT_BORDER)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:cantSplit")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def keep_paragraph_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "目录将在 Word 中自动更新"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_inline_markup(paragraph, text, size=10.5):
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    math_replacements = {
        "$v_i$": "vᵢ",
        "$g$": "g",
        "$w_{i,g}$": "w(i,g)",
        "$V$": "V",
    }
    for source, replacement in math_replacements.items():
        text = text.replace(source, replacement)
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, east_asia="等线", latin="Consolas", size=max(8.5, size - 1), color="243B53")
            set_cell = run._element.get_or_add_rPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "F2F2F2")
            set_cell.append(shd)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True)
        else:
            run = paragraph.add_run(part.replace("—", "至"))
            set_run_font(run, size=size)


def style_body_paragraph(paragraph, first_line=True):
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.line_spacing = 1.45
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(5)
    if first_line:
        fmt.first_line_indent = Pt(21)
    fmt.widow_control = True


def add_table(doc, rows):
    headers = rows[0]
    body_rows = rows[1:]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    page_width_cm = 16.4
    lengths = [max(5, max(len(str(row[i])) if i < len(row) else 0 for row in rows)) for i in range(len(headers))]
    total = sum(min(value, 35) for value in lengths)
    widths = [page_width_cm * min(value, 35) / total for value in lengths]
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        cell.width = Cm(widths[i])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, east_asia="黑体", latin="Arial", size=9, bold=True, color="FFFFFF")
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for row_index, values in enumerate(body_rows, start=1):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        fill = "FFFFFF" if row_index % 2 else PALE_BLUE
        for i, value in enumerate(values):
            cell = cells[i]
            cell.width = Cm(widths[i])
            set_cell_shading(cell, fill)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(0)
            numeric = bool(re.fullmatch(r"[\d,.%–\- 至]+", value.strip()))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if numeric or len(value) <= 14 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline_markup(p, value, size=8.6)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_picture(doc, path, max_width=5.85, max_height=7.0):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(2)
    with Image.open(path) as image:
        width_px, height_px = image.size
    aspect = width_px / height_px
    render_width = min(max_width, max_height * aspect)
    paragraph.add_run().add_picture(str(path), width=Inches(render_width))
    paragraph.paragraph_format.keep_with_next = True


def add_native_flow(doc, labels):
    table = doc.add_table(rows=1, cols=len(labels))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for i, label in enumerate(labels):
        cell = table.cell(0, i)
        cell.width = Cm(16.4 / len(labels))
        set_cell_shading(cell, PALE_BLUE if i % 2 == 0 else "F7F2EA")
        set_cell_margins(cell, top=130, bottom=130, start=70, end=70)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label + ("\n→" if i < len(labels) - 1 else ""))
        set_run_font(r, east_asia="黑体", latin="Arial", size=8.2, bold=True, color=NAVY)
    keep_paragraph_with_next(table.cell(0, 0).paragraphs[0])


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, east_asia="宋体", latin="Times New Roman", size=9, color=MUTED)


def add_equation(doc, raw):
    if "HHI" in raw:
        normalized = "HHI = 10000 × Σ[g] (E(g) / V)²    (5-2)"
    elif "Delta" in raw:
        normalized = "Δv(i) = V × [w(i,target) − w(i,current)]    (5-3)"
    elif "E_g" in raw:
        normalized = "E(g) = Σ[i] v(i) × w(i,g),    p(g) = E(g) / V × 100%    (5-1)"
    else:
        normalized = re.sub(r"\\[a-zA-Z]+", "", raw).replace("{", "").replace("}", "")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(normalized)
    set_run_font(r, east_asia="Cambria Math", latin="Cambria Math", size=11)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BODY)

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(6 if name == "Heading 1" else 4)
        style.paragraph_format.keep_with_next = True
        if name == "Heading 1":
            style.paragraph_format.page_break_before = True

    caption = doc.styles["Caption"]
    caption.font.name = "Times New Roman"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string(MUTED)


def add_cover(doc):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(105)
    p.paragraph_format.space_after = Pt(24)
    p.add_run("Prism 个性化证券投顾智能体系统项目详细方案")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(42)
    r = p.add_run("基于同花顺问财 SkillHub 的个性化证券投顾智能体系统设计")
    set_run_font(r, east_asia="黑体", latin="Arial", size=13, color="333333")

    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    meta_data = [
        ("项目名称", "Prism 个性化证券投顾智能体系统"),
        ("参赛方向", "人工智能  智能体"),
        ("核心能力", "画像约束  多智能体研究  确定性计算  双重审查"),
        ("文档日期", "2026 年 9 月"),
    ]
    for i, (key, value) in enumerate(meta_data):
        for j, text in enumerate((key, value)):
            cell = meta.cell(i, j)
            cell.width = Cm(3.2 if j == 0 else 10.8)
            set_cell_shading(cell, PALE_BLUE if j == 0 else "FFFFFF")
            set_cell_margins(cell, top=110, bottom=110, start=140, end=140)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            set_run_font(r, east_asia="黑体" if j == 0 else "宋体", size=10, bold=j == 0)
    set_table_borders(meta)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(34)
    p.paragraph_format.line_spacing = 1.55
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(
        "本方案说明 Prism 如何把投资者画像、专业研究、金融计算、证据追溯和风险合规审查连接为一条可验证的投顾决策链。"
        "系统覆盖大盘、行业、个股、基金、可转债和组合优化场景，所有数值由确定性程序计算，所有建议在输出前经过独立风险与合规审查。"
    )
    set_run_font(r, size=11)
    doc.add_page_break()


def add_toc(doc):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = False
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("目录")
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(8)
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = intro.add_run("项目详细方案章节索引")
    set_run_font(run, east_asia="宋体", size=9.5, color=MUTED)
    for title, page in TOC_ENTRIES:
        toc = doc.add_paragraph()
        toc.paragraph_format.left_indent = Cm(0.6)
        toc.paragraph_format.right_indent = Cm(0.6)
        toc.paragraph_format.space_before = Pt(2)
        toc.paragraph_format.space_after = Pt(4)
        toc.paragraph_format.tab_stops.add_tab_stop(Cm(15.1), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        run = toc.add_run(f"{title}\t{page}")
        set_run_font(run, east_asia="宋体", size=10.5, bold=title.startswith(("第一章", "第五章", "第六章", "第七章")))
    doc.add_page_break()


def add_header_footer(section):
    section.different_first_page_header_footer = True
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Prism 个性化证券投顾智能体系统项目详细方案")
    set_run_font(r, east_asia="宋体", size=8.5, color="777777")
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("第 ")
    set_run_font(r, size=9, color="777777")
    add_field(p, "PAGE")
    r = p.add_run(" 页")
    set_run_font(r, size=9, color="777777")


def parse_markdown(doc):
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    start = next(i for i, line in enumerate(lines) if line.strip() == "## 第一章 项目概况")
    i = start
    pending_mermaid = False
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if stripped.startswith("<a id="):
            i += 1
            continue
        if stripped.startswith("```mermaid"):
            pending_mermaid = True
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                i += 1
            i += 1
            continue
        if stripped == "$$":
            equation = []
            i += 1
            while i < len(lines) and lines[i].strip() != "$$":
                equation.append(lines[i].strip())
                i += 1
            add_equation(doc, " ".join(equation))
            i += 1
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            parsed = [[part.strip() for part in row.strip("|").split("|")] for row in table_lines]
            if len(parsed) >= 2 and all(re.fullmatch(r":?-+:?", cell.replace(" ", "")) for cell in parsed[1]):
                parsed.pop(1)
            add_table(doc, parsed)
            continue
        if stripped.startswith("!["):
            match = re.match(r"!\[(.*?)\]\((.*?)\)", stripped)
            if match:
                path = (SOURCE.parent / match.group(2)).resolve()
                if path.exists():
                    add_picture(doc, path)
            i += 1
            continue
        if stripped.startswith("图 "):
            key_match = re.match(r"(图 \d+-\d+)", stripped)
            key = key_match.group(1) if key_match else ""
            if pending_mermaid:
                if key in FIGURE_MAP and FIGURE_MAP[key].exists():
                    add_picture(doc, FIGURE_MAP[key])
                elif key in FLOW_LABELS:
                    add_native_flow(doc, FLOW_LABELS[key])
                pending_mermaid = False
            add_caption(doc, stripped)
            i += 1
            continue
        if stripped.startswith("表 "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(3)
            keep_paragraph_with_next(p)
            r = p.add_run(stripped)
            set_run_font(r, east_asia="黑体", size=9.5, bold=True)
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=1)
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
            i += 1
            continue
        if stripped.startswith("#### "):
            doc.add_heading(stripped[5:], level=3)
            i += 1
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if numbered:
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.35
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.left_indent = Pt(21)
            p.paragraph_format.first_line_indent = Pt(-21)
            add_inline_markup(p, f"{numbered.group(1)}.  {numbered.group(2)}")
            i += 1
            continue
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing = 1.35
            p.paragraph_format.space_after = Pt(3)
            add_inline_markup(p, stripped[2:])
            i += 1
            continue
        if stripped.startswith("> "):
            p = doc.add_paragraph()
            style_body_paragraph(p, first_line=False)
            add_inline_markup(p, stripped[2:])
            i += 1
            continue
        if not stripped:
            i += 1
            continue
        p = doc.add_paragraph()
        style_body_paragraph(p)
        add_inline_markup(p, stripped)
        i += 1


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.25)
    section.bottom_margin = Cm(2.15)
    section.left_margin = Cm(2.25)
    section.right_margin = Cm(2.25)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)
    configure_styles(doc)
    add_header_footer(section)
    add_cover(doc)
    add_toc(doc)
    parse_markdown(doc)
    settings = doc.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)
    doc.core_properties.title = "Prism 个性化证券投顾智能体系统项目详细方案"
    doc.core_properties.subject = "同花顺 A18 个性化证券投顾智能体系统竞赛项目详细方案"
    doc.core_properties.author = "Prism 项目团队"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
